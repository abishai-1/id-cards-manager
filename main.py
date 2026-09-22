import os
import io
import json
import uuid
import datetime
import base64
import binascii
import re
from flask import Flask, request, jsonify, send_from_directory, redirect
import pymysql
import pymysql.cursors
from dotenv import load_dotenv
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.utils import secure_filename

# Load Environment Variables
# .env
load_dotenv()

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", uuid.uuid4().hex)
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)
app.config["MAX_CONTENT_LENGTH"] = int(os.getenv("MAX_UPLOAD_MB", "50")) * 1024 * 1024

SITE_DOMAIN = os.getenv("SITE_DOMAIN", "con.satishadagency.in").strip().lower()
SITE_URL = os.getenv("SITE_URL", f"https://{SITE_DOMAIN}").rstrip("/")
ENFORCE_CANONICAL_DOMAIN = os.getenv("ENFORCE_CANONICAL_DOMAIN", "false").lower() == "true"

# DB Configurations
DB_HOST = os.getenv("DB_HOST", "172.18.0.1")
DB_PORT = int(os.getenv("DB_PORT", 3306))
DB_USER = os.getenv("DB_USER", "")
DB_PASSWORD = os.getenv("DB_PASSWORD", "")
DB_NAME = os.getenv("DB_NAME", "s52_satish")

ADMIN_USERNAME = os.getenv("ADMIN_USERNAME", "satish")
ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD", "")
ALLOWED_EXCEL_EXTENSIONS = {"xlsx"}
ALLOWED_IMAGE_EXTENSIONS = {"jpg", "jpeg", "png", "gif", "webp"}
ALLOWED_IMAGE_MIME_EXTENSIONS = {"png", "jpg", "jpeg", "gif", "webp"}
PUBLIC_STATIC_FILES = {"index.html", "app.html", "contact.html", "styles.css", "app.js", "favicon.jpeg"}
UPLOAD_FOLDER = os.getenv("UPLOAD_FOLDER", "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

CATALOG_ITEM_KEYS = {
    "id-cards",
    "belts-ties",
    "school-diaries",
    "progress-reports",
    "book-cover-pages",
    "student-files",
    "student-file-folders",
    "answer-sheets",
    "note-books",
    "cloth-ribbon-badges",
    "pocket-pvc-badges",
    "prize-medals",
    "prize-shields",
    "t-shirts-tracks",
    "school-socks",
    "school-shoes",
    "caps",
    "brochures-posters",
    "display-boards",
    "non-woven-bags",
    "school-bags",
    "and-more",
}

STUDENT_SORT_COLUMNS = {
    "id": "id",
    "roll": "roll_number",
    "roll_number": "roll_number",
    "photo": "photo_data",
    "name": "name",
    "class": "class",
    "section": "section",
    "father": "father_name",
    "father_name": "father_name",
    "mother": "mother_name",
    "mother_name": "mother_name",
    "dob": "dob",
    "blood": "blood_group",
    "blood_group": "blood_group",
    "contact": "contact",
    "address": "address",
    "status": "status",
}


def get_file_extension(filename):
    return filename.rsplit(".", 1)[1].lower() if filename and "." in filename else ""


def normalize_identifier(value, strip_symbols=False):
    """Build a stable comparison key for spreadsheet photo IDs and filenames."""
    text = os.path.basename(str(value or "").strip().lower())
    text = os.path.splitext(text)[0]
    text = re.sub(r"\s+", "", text)
    if strip_symbols:
        text = re.sub(r"[^a-z0-9]+", "", text)
    return text


def identifier_keys(value):
    """Return increasingly forgiving keys while keeping ambiguous matches detectable."""
    exact_key = normalize_identifier(value)
    canonical_key = normalize_identifier(value, strip_symbols=True)
    keys = [key for key in (exact_key, canonical_key) if key]
    if canonical_key.isdigit():
        keys.append(canonical_key.lstrip("0") or "0")
    return list(dict.fromkeys(keys))


def loose_text_matches(stored_value, entered_value):
    stored = str(stored_value or "").strip()
    entered = str(entered_value or "").strip()
    if not stored:
        return True
    if not entered:
        return False
    return normalize_identifier(stored, strip_symbols=True) == normalize_identifier(entered, strip_symbols=True)


def optional_identifier_matches(stored_value, entered_value):
    stored = str(stored_value or "").strip()
    entered = str(entered_value or "").strip()
    if not stored:
        return True
    if not entered:
        return False
    return bool(set(identifier_keys(entered)).intersection(identifier_keys(stored)))


def normalize_header(value):
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def looks_like_photo_reference(value):
    ext = get_file_extension(str(value or "").strip())
    return ext in ALLOWED_IMAGE_EXTENSIONS


def find_header_index(headers, aliases, excluded_indexes=None):
    excluded_indexes = set(excluded_indexes or [])
    normalized_headers = [normalize_header(header) for header in headers]
    normalized_aliases = [normalize_header(alias) for alias in aliases]

    for alias in normalized_aliases:
        for idx, header in enumerate(normalized_headers):
            if idx not in excluded_indexes and header == alias:
                return idx

    for alias in normalized_aliases:
        if len(alias) < 4:
            continue
        for idx, header in enumerate(normalized_headers):
            if idx not in excluded_indexes and alias in header:
                return idx
    return -1


def add_photo_candidate(candidate_map, value, entry):
    for key in identifier_keys(value):
        candidate_map.setdefault(key, []).append(entry)


def find_unique_photo_match(candidate_map, *values):
    """Match only a single uploaded file; collisions stay unmatched for manual review."""
    for value in values:
        unique_candidates = {}
        for key in identifier_keys(value):
            for candidate in candidate_map.get(key, []):
                unique_candidates[candidate["path"]] = candidate
        if len(unique_candidates) == 1:
            return next(iter(unique_candidates.values()))
    return None


def generate_access_code():
    return f"SAT-{uuid.uuid4().hex[:8].upper()}"


def parse_bool_arg(name, default=False):
    value = request.args.get(name)
    if value is None:
        return default
    return str(value).strip().lower() in {"1", "true", "yes", "on"}


def looks_like_image(file_storage):
    pos = file_storage.stream.tell()
    header = file_storage.stream.read(16)
    file_storage.stream.seek(pos)
    return (
        header.startswith(b"\xff\xd8\xff")
        or header.startswith(b"\x89PNG\r\n\x1a\n")
        or header.startswith(b"GIF87a")
        or header.startswith(b"GIF89a")
        or (header.startswith(b"RIFF") and header[8:12] == b"WEBP")
    )


def looks_like_xlsx(file_storage):
    pos = file_storage.stream.tell()
    header = file_storage.stream.read(4)
    file_storage.stream.seek(pos)
    return header.startswith(b"PK\x03\x04")


def save_upload_file(file_storage, *parts):
    ext = get_file_extension(file_storage.filename)
    safe_parts = [secure_filename(str(part)) for part in parts if str(part).strip()]
    upload_dir = os.path.join(UPLOAD_FOLDER, *safe_parts)
    os.makedirs(upload_dir, exist_ok=True)

    original_name = secure_filename(file_storage.filename) or f"upload.{ext}"
    stem = os.path.splitext(original_name)[0] or "upload"
    filename = f"{stem}-{uuid.uuid4().hex[:12]}.{ext}"
    file_path = os.path.join(upload_dir, filename)
    file_storage.save(file_path)
    return "/uploads/" + os.path.relpath(file_path, UPLOAD_FOLDER).replace("\\", "/")


def save_data_url(data_url, filename, *parts):
    if not data_url or not str(data_url).startswith("data:"):
        return data_url

    try:
        header, encoded = data_url.split(",", 1)
    except ValueError as exc:
        raise ValueError("Invalid image data.") from exc

    mime = header.split(";", 1)[0].replace("data:", "")
    ext = mime.split("/", 1)[1].lower() if "/" in mime else "jpg"
    if ext == "jpeg":
        ext = "jpg"
    if ext not in ALLOWED_IMAGE_MIME_EXTENSIONS:
        raise ValueError("Unsupported image data type.")
    if ";base64" not in header:
        raise ValueError("Image data must be base64 encoded.")

    safe_parts = [secure_filename(str(part)) for part in parts if str(part).strip()]
    upload_dir = os.path.join(UPLOAD_FOLDER, *safe_parts)
    os.makedirs(upload_dir, exist_ok=True)

    safe_filename = secure_filename(filename) or f"image.{ext}"
    stem = os.path.splitext(safe_filename)[0] or "image"
    file_path = os.path.join(upload_dir, f"{stem}-{uuid.uuid4().hex[:12]}.{ext}")
    with open(file_path, "wb") as image_file:
        try:
            image_file.write(base64.b64decode(encoded, validate=True))
        except binascii.Error as exc:
            raise ValueError("Invalid image data.") from exc

    return "/uploads/" + os.path.relpath(file_path, UPLOAD_FOLDER).replace("\\", "/")


def split_class_section(value):
    text = str(value or "").strip()
    if not text:
        return "", ""

    for separator in (" - ", "-", "/", "|", ","):
        if separator in text:
            class_part, section_part = text.split(separator, 1)
            return class_part.strip(), section_part.strip()

    parts = text.split()
    if len(parts) >= 2:
        return parts[0].strip(), " ".join(parts[1:]).strip()

    return text, ""

def get_db_connection(select_db=True):
    """
    Establishes a connection to the MySQL server.
    """
    return pymysql.connect(
        host=DB_HOST,
        port=DB_PORT,
        user=DB_USER,
        password=DB_PASSWORD,
        database=DB_NAME if select_db else None,
        charset='utf8mb4',
        cursorclass=pymysql.cursors.DictCursor
    )

def init_db():
    """
    Initializes the database and creates necessary tables automatically.
    """
    print("Initializing database...")
    # Step 1: Create Database if not exists
    try:
        conn = get_db_connection(select_db=False)
        with conn.cursor() as cursor:
            cursor.execute(f"CREATE DATABASE IF NOT EXISTS {DB_NAME}")
        conn.close()
        print(f"Database '{DB_NAME}' verified/created.")
    except Exception as e:
        print("Warning: Could not check/create database from root connection. Proceeding directly.", e)

    # Step 2: Establish connection and create tables
    conn = get_db_connection(select_db=True)
    try:
        with conn.cursor() as cursor:
            # 1. Projects Table
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS projects (
                    id INT AUTO_INCREMENT PRIMARY KEY,
                    name VARCHAR(255) NOT NULL UNIQUE,
                    access_code VARCHAR(64) UNIQUE
                ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;
            """)

            # 2. Students Table
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS students (
                    id VARCHAR(255) NOT NULL,
                    project_id INT NOT NULL,
                    photo_id VARCHAR(255),
                    roll_number VARCHAR(50),
                    name VARCHAR(255) NOT NULL,
                    class VARCHAR(50),
                    section VARCHAR(50),
                    father_name VARCHAR(255),
                    mother_name VARCHAR(255),
                    contact VARCHAR(50),
                    dob VARCHAR(50),
                    blood_group VARCHAR(50),
                    address TEXT,
                    correction_note TEXT,
                    status VARCHAR(50) DEFAULT 'pending',
                    edit_count INT DEFAULT 0,
                    photo_data LONGTEXT,
                    PRIMARY KEY (project_id, id),
                    FOREIGN KEY (project_id) REFERENCES projects(id) ON DELETE CASCADE
                ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;
            """)

            # 3. Project Configs (ID Card Designer detail storage)
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS project_configs (
                    project_id INT PRIMARY KEY,
                    frame_data LONGTEXT,
                    config_json TEXT,
                    FOREIGN KEY (project_id) REFERENCES projects(id) ON DELETE CASCADE
                ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;
            """)

            # 4. Sessions Table
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS sessions (
                    token VARCHAR(255) PRIMARY KEY,
                    role VARCHAR(50) NOT NULL,
                    project_id INT NULL,
                    student_id VARCHAR(255) NULL,
                    expires_at DATETIME NOT NULL
                ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;
            """)

            # 5. Contact Messages Table
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS contact_messages (
                    id INT AUTO_INCREMENT PRIMARY KEY,
                    name VARCHAR(255) NOT NULL,
                    email VARCHAR(255) NOT NULL,
                    message TEXT NOT NULL,
                    created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP
                ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;
            """)

            # 6. Public catalog image overrides managed by admins
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS catalog_items (
                    item_key VARCHAR(100) PRIMARY KEY,
                    images_json LONGTEXT NOT NULL,
                    updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP
                ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;
            """)

            cursor.execute("SHOW COLUMNS FROM students LIKE 'correction_note'")
            if not cursor.fetchone():
                cursor.execute("ALTER TABLE students ADD COLUMN correction_note TEXT AFTER address")
            cursor.execute("SHOW COLUMNS FROM students LIKE 'photo_id'")
            if not cursor.fetchone():
                cursor.execute("ALTER TABLE students ADD COLUMN photo_id VARCHAR(255) AFTER project_id")
            cursor.execute("UPDATE students SET photo_id = id WHERE photo_id IS NULL OR photo_id = ''")
            cursor.execute("SHOW COLUMNS FROM projects LIKE 'access_code'")
            if not cursor.fetchone():
                cursor.execute("ALTER TABLE projects ADD COLUMN access_code VARCHAR(64) UNIQUE AFTER name")
            cursor.execute("SELECT id FROM projects WHERE access_code IS NULL OR access_code = ''")
            for project in cursor.fetchall():
                cursor.execute(
                    "UPDATE projects SET access_code = %s WHERE id = %s",
                    (generate_access_code(), project["id"])
                )
        conn.commit()
        print("Database schema successfully verified/initialized.")
    except Exception as e:
        print("Error initializing database schema:", e)
        raise e
    finally:
        conn.close()

# Initialize DB on server start. Tests can opt out before importing this module.
if os.getenv("SKIP_DB_INIT", "false").lower() != "true":
    try:
        init_db()
    except Exception as err:
        print("Database initialization failed. Please verify credentials in your .env configuration.", err)

# AUTHENTICATION DECORATOR & HELPERS
def get_current_session():
    """
    Checks the authorization header and returns session context if valid.
    """
    auth_header = request.headers.get("Authorization")
    if not auth_header:
        return None
    
    token = auth_header
    if auth_header.startswith("Bearer "):
        token = auth_header[7:]

    conn = get_db_connection()
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                "SELECT * FROM sessions WHERE token = %s AND expires_at > NOW()",
                (token,)
            )
            return cursor.fetchone()
    except Exception as e:
        print("Session check failure:", e)
        return None
    finally:
        conn.close()

# ROUTING & REST API
@app.before_request
def redirect_to_canonical_domain():
    if not ENFORCE_CANONICAL_DOMAIN or not SITE_DOMAIN:
        return None

    request_host = request.host.split(":", 1)[0].lower()
    if request_host == SITE_DOMAIN:
        return None

    return redirect(f"{SITE_URL}{request.full_path}".rstrip("?"), code=301)

@app.route('/')
def serve_index():
    return send_from_directory('.', 'index.html')

@app.route('/index.html')
def serve_index_explicit():
    return send_from_directory('.', 'index.html')

@app.route('/app.html')
def serve_app():
    return send_from_directory('.', 'app.html')

@app.route('/auth/login')
def serve_login_alias():
    return redirect('/app.html', code=302)

@app.route('/contact.html')
def serve_contact():
    return send_from_directory('.', 'contact.html')

@app.route('/uploads/<path:filename>')
def serve_upload(filename):
    return send_from_directory(UPLOAD_FOLDER, filename)

@app.errorhandler(404)
def not_found(_error):
    if request.path.startswith('/api/'):
        return jsonify({"success": False, "message": "API endpoint not found"}), 404
    return "Page Not Found", 404

@app.route('/<path:path>')
def serve_static(path):
    if path in PUBLIC_STATIC_FILES:
        return send_from_directory('.', path)
    if path.startswith('api/'):
        return jsonify({"success": False, "message": "API endpoint not found"}), 404
    return "Page Not Found", 404

# API: Auth login
@app.route('/api/auth/login', methods=['POST'])
def api_login():
    data = request.get_json(silent=True) or {}
    role = data.get("role")

    if role == "admin" and not ADMIN_PASSWORD:
        return jsonify({"success": False, "message": "Admin login is not configured."}), 503

    conn = None
    try:
        conn = get_db_connection()
        with conn.cursor() as cursor:
            # 1. Admin login
            if role == "admin":
                username = data.get("username", "").strip()
                password = data.get("password", "").strip()
                if username.lower() == ADMIN_USERNAME.lower() and password == ADMIN_PASSWORD:
                    token = str(uuid.uuid4())
                    expires = datetime.datetime.now() + datetime.timedelta(days=1)
                    cursor.execute(
                        "INSERT INTO sessions (token, role, expires_at) VALUES (%s, 'admin', %s)",
                        (token, expires)
                    )
                    conn.commit()
                    return jsonify({
                        "success": True,
                        "token": token,
                        "role": "admin",
                        "userContext": {"name": "Satish (Admin)"}
                    })
                return jsonify({"success": False, "message": "Invalid Admin credentials!"}), 401
            
            # 2. Staff login (using the project access code generated for the school)
            elif role == "staff":
                code = data.get("code", "").strip()
                cursor.execute("SELECT * FROM projects WHERE access_code = %s", (code,))
                project = cursor.fetchone()
                if project:
                    token = str(uuid.uuid4())
                    expires = datetime.datetime.now() + datetime.timedelta(days=1)
                    cursor.execute(
                        "INSERT INTO sessions (token, role, project_id, expires_at) VALUES (%s, 'staff', %s, %s)",
                        (token, project["id"], expires)
                    )
                    conn.commit()
                    return jsonify({
                        "success": True,
                        "token": token,
                        "role": "staff",
                        "userContext": {
                            "project_id": project["id"],
                            "project_name": project["name"]
                        }
                    })
                return jsonify({"success": False, "message": "Invalid School/Access code!"}), 401

            # 3. Student login
            elif role == "student":
                school = data.get("school", "").strip()
                class_value = data.get("class", "").strip()
                section = data.get("section", "").strip()
                roll_number = data.get("roll_number", "").strip()
                photo_id = data.get("photo_id", "").strip()
                if not school or not any((class_value, section, roll_number, photo_id)):
                    return jsonify({
                        "success": False,
                        "message": "School and at least one student detail are required."
                    }), 400

                # Scope to the selected school first. Optional student details are compared only
                # when the master database has those values for a candidate record.
                cursor.execute("""
                    SELECT s.*
                    FROM students s
                    INNER JOIN projects p ON p.id = s.project_id
                    WHERE LOWER(TRIM(p.name)) = LOWER(TRIM(%s))
                      AND s.status <> 'archived'
                """, (school,))
                students = [
                    candidate for candidate in cursor.fetchall()
                    if loose_text_matches(candidate.get("class"), class_value)
                    and loose_text_matches(candidate.get("section"), section)
                    and loose_text_matches(candidate.get("roll_number"), roll_number)
                    and optional_identifier_matches(candidate.get("photo_id"), photo_id)
                ]
                if len(students) > 1:
                    return jsonify({
                        "success": False,
                        "message": "We could not uniquely identify your profile. Please retry with more details or contact your school staff."
                    }), 409
                student = students[0] if students else None
                if student:
                    token = str(uuid.uuid4())
                    expires = datetime.datetime.now() + datetime.timedelta(days=1)
                    cursor.execute(
                        "INSERT INTO sessions (token, role, project_id, student_id, expires_at) VALUES (%s, 'student', %s, %s, %s)",
                        (token, student["project_id"], student["id"], expires)
                    )
                    conn.commit()
                    return jsonify({
                        "success": True,
                        "token": token,
                        "role": "student",
                        "userContext": {
                            "id": student["id"],
                            "photo_id": student.get("photo_id") or student["id"],
                            "name": student["name"],
                            "project_id": student["project_id"]
                        }
                    })
                return jsonify({
                    "success": False,
                    "message": "No profile was found with the available details. Please retry or contact your school staff."
                }), 404
            
            return jsonify({"success": False, "message": "Invalid role selected."}), 400
    except Exception as e:
        print("Login processing error:", e)
        return jsonify({
            "success": False,
            "message": "Database unavailable. Please verify the server database configuration."
        }), 503
    finally:
        if conn:
            conn.close()

# API: Log out
@app.route('/api/auth/logout', methods=['POST'])
def api_logout():
    auth_header = request.headers.get("Authorization")
    if auth_header:
        token = auth_header
        if auth_header.startswith("Bearer "):
            token = auth_header[7:]
        
        conn = get_db_connection()
        try:
            with conn.cursor() as cursor:
                cursor.execute("DELETE FROM sessions WHERE token = %s", (token,))
            conn.commit()
        except Exception as e:
            print("Logout error:", e)
        finally:
            conn.close()
    return jsonify({"success": True})

@app.route('/api/contact', methods=['GET', 'POST'])
def api_contact():
    if request.method == 'GET':
        session = get_current_session()
        if not session:
            return jsonify({"success": False, "message": "Session expired or invalid"}), 401
        if session["role"] != "admin":
            return jsonify({"success": False, "message": "Unauthorized"}), 403

        limit = max(1, min(request.args.get("limit", default=100, type=int), 500))
        conn = get_db_connection()
        try:
            with conn.cursor() as cursor:
                cursor.execute(
                    "SELECT id, name, email, message, created_at FROM contact_messages ORDER BY created_at DESC LIMIT %s",
                    (limit,)
                )
                messages = cursor.fetchall()
            return jsonify({"success": True, "messages": messages})
        except Exception as e:
            print("Contact inbox API error:", e)
            return jsonify({"success": False, "message": "Could not load contact messages."}), 500
        finally:
            conn.close()

    data = request.get_json(silent=True) or {}
    name = data.get("name", "").strip()
    email = data.get("email", "").strip()
    message = data.get("message", "").strip()

    if not name or not email or not message:
        return jsonify({"success": False, "message": "Name, email, and message are required."}), 400
    if "@" not in email or "." not in email:
        return jsonify({"success": False, "message": "Please enter a valid email address."}), 400
    if len(message) > 5000:
        return jsonify({"success": False, "message": "Message is too long."}), 400

    conn = get_db_connection()
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                "INSERT INTO contact_messages (name, email, message) VALUES (%s, %s, %s)",
                (name, email, message)
            )
        conn.commit()
        return jsonify({"success": True, "message": "Message sent successfully."})
    except Exception as e:
        print("Contact API error:", e)
        return jsonify({"success": False, "message": "Could not send message."}), 500
    finally:
        conn.close()


@app.route('/api/contact/<int:message_id>', methods=['DELETE'])
def api_delete_contact_message(message_id):
    session = get_current_session()
    if not session:
        return jsonify({"success": False, "message": "Session expired or invalid"}), 401
    if session["role"] != "admin":
        return jsonify({"success": False, "message": "Unauthorized"}), 403

    conn = get_db_connection()
    try:
        with conn.cursor() as cursor:
            cursor.execute("DELETE FROM contact_messages WHERE id = %s", (message_id,))
        conn.commit()
        return jsonify({"success": True, "message": "Message deleted successfully"})
    except Exception as e:
        print("Delete contact message error:", e)
        return jsonify({"success": False, "message": "Could not delete message."}), 500
    finally:
        conn.close()


# API: Public catalog images and admin-only image replacement
@app.route('/api/catalog', methods=['GET'])
def api_catalog():
    conn = None
    try:
        conn = get_db_connection()
        with conn.cursor() as cursor:
            cursor.execute("SELECT item_key, images_json FROM catalog_items")
            rows = cursor.fetchall()

        catalog = {}
        for row in rows:
            item_key = row.get("item_key")
            if item_key not in CATALOG_ITEM_KEYS:
                continue
            try:
                images = json.loads(row.get("images_json") or "[]")
            except (TypeError, json.JSONDecodeError):
                continue
            if 4 <= len(images) <= 6 and all(isinstance(image, str) and image for image in images):
                catalog[item_key] = images
        return jsonify({"success": True, "catalog": catalog})
    except Exception as e:
        # The public homepage keeps its built-in gallery if the database is temporarily unavailable.
        print("Catalog API read error:", e)
        return jsonify({"success": True, "catalog": {}})
    finally:
        if conn:
            conn.close()


@app.route('/api/catalog/<item_key>', methods=['PUT'])
def api_update_catalog_item(item_key):
    session = get_current_session()
    if not session:
        return jsonify({"success": False, "message": "Session expired or invalid"}), 401
    if session["role"] != "admin":
        return jsonify({"success": False, "message": "Unauthorized"}), 403
    if item_key not in CATALOG_ITEM_KEYS:
        return jsonify({"success": False, "message": "Catalog item not found"}), 404

    try:
        uploaded_images = request.files.getlist("images")
        if uploaded_images:
            if not (4 <= len(uploaded_images) <= 6):
                return jsonify({"success": False, "message": "Please upload 4 to 6 catalog images."}), 400
            saved_images = []
            for image_file in uploaded_images:
                ext = get_file_extension(image_file.filename)
                if ext not in ALLOWED_IMAGE_EXTENSIONS:
                    return jsonify({"success": False, "message": f"Unsupported image type: {image_file.filename}"}), 400
                if not looks_like_image(image_file):
                    return jsonify({"success": False, "message": f"Invalid image file: {image_file.filename}"}), 400
                saved_images.append(save_upload_file(image_file, "catalog", item_key))
        else:
            data = request.get_json(silent=True) or {}
            images = data.get("images")
            if not isinstance(images, list) or not (4 <= len(images) <= 6):
                return jsonify({"success": False, "message": "Please upload 4 to 6 catalog images."}), 400

            existing_prefix = f"/uploads/catalog/{secure_filename(item_key)}/"
            for image in images:
                if not isinstance(image, str) or not image:
                    return jsonify({"success": False, "message": "Each catalog image is required."}), 400
                if not image.startswith("data:") and not image.startswith(existing_prefix):
                    return jsonify({"success": False, "message": "Catalog images must be uploaded through the admin editor."}), 400

            saved_images = [
                save_data_url(image, f"{item_key}-{index}.jpg", "catalog", item_key)
                for index, image in enumerate(images, start=1)
            ]
    except ValueError as err:
        return jsonify({"success": False, "message": str(err)}), 400

    conn = get_db_connection()
    try:
        with conn.cursor() as cursor:
            cursor.execute("""
                INSERT INTO catalog_items (item_key, images_json)
                VALUES (%s, %s)
                ON DUPLICATE KEY UPDATE images_json = VALUES(images_json)
            """, (item_key, json.dumps(saved_images)))
        conn.commit()
        return jsonify({
            "success": True,
            "message": "Catalog images updated successfully.",
            "item_key": item_key,
            "images": saved_images,
        })
    except Exception as e:
        print("Catalog API save error:", e)
        return jsonify({"success": False, "message": "Could not save catalog images."}), 500
    finally:
        conn.close()


# API: Projects Management (Admin only)
@app.route('/api/projects', methods=['GET', 'POST'])
def api_projects():
    session = get_current_session()
    if not session:
        return jsonify({"success": False, "message": "Session expired or invalid"}), 401
    if session["role"] != "admin":
        return jsonify({"success": False, "message": "Unauthorized"}), 403

    conn = get_db_connection()
    try:
        with conn.cursor() as cursor:
            if request.method == 'GET':
                cursor.execute("SELECT id, name, access_code FROM projects ORDER BY name ASC")
                projects = cursor.fetchall()
                return jsonify({"success": True, "projects": projects})
            
            elif request.method == 'POST':
                data = request.get_json(silent=True) or {}
                name = data.get("name", "").strip()
                if not name:
                    return jsonify({"success": False, "message": "Project name required"}), 400
                
                try:
                    access_code = generate_access_code()
                    cursor.execute("INSERT INTO projects (name, access_code) VALUES (%s, %s)", (name, access_code))
                    conn.commit()
                    cursor.execute("SELECT * FROM projects WHERE name = %s", (name,))
                    new_proj = cursor.fetchone()
                    return jsonify({"success": True, "project": new_proj})
                except pymysql.err.IntegrityError:
                    return jsonify({"success": False, "message": "Project with this name already exists!"}), 409
    except Exception as e:
        print("Projects API error:", e)
        return jsonify({"success": False, "message": "Internal Server Error"}), 500
    finally:
        conn.close()

@app.route('/api/projects/<int:project_id>', methods=['DELETE'])
def api_delete_project(project_id):
    session = get_current_session()
    if not session:
        return jsonify({"success": False, "message": "Session expired or invalid"}), 401
    if session["role"] != "admin":
        return jsonify({"success": False, "message": "Unauthorized"}), 403

    conn = get_db_connection()
    try:
        with conn.cursor() as cursor:
            cursor.execute("DELETE FROM projects WHERE id = %s", (project_id,))
            conn.commit()
            return jsonify({"success": True, "message": "Project deleted successfully"})
    except Exception as e:
        print("Project deletion error:", e)
        return jsonify({"success": False, "message": "Internal Server Error"}), 500
    finally:
        conn.close()


@app.route('/api/projects/<int:project_id>/access-code/regenerate', methods=['POST'])
def api_regenerate_project_access_code(project_id):
    session = get_current_session()
    if not session:
        return jsonify({"success": False, "message": "Session expired or invalid"}), 401
    if session["role"] != "admin":
        return jsonify({"success": False, "message": "Unauthorized"}), 403

    new_code = generate_access_code()
    conn = get_db_connection()
    try:
        with conn.cursor() as cursor:
            cursor.execute("UPDATE projects SET access_code = %s WHERE id = %s", (new_code, project_id))
            conn.commit()
            cursor.execute("SELECT id, name, access_code FROM projects WHERE id = %s", (project_id,))
            project = cursor.fetchone()
            if not project:
                return jsonify({"success": False, "message": "Project not found"}), 404
            return jsonify({"success": True, "project": project})
    except Exception as e:
        print("Access code regeneration error:", e)
        return jsonify({"success": False, "message": "Internal Server Error"}), 500
    finally:
        conn.close()


# API: ID card config management (Admin only)
@app.route('/api/config', methods=['GET', 'POST'])
def api_config():
    session = get_current_session()
    if not session:
        return jsonify({"success": False, "message": "Session expired or invalid"}), 401
    if session["role"] != "admin":
        return jsonify({"success": False, "message": "Unauthorized"}), 403

    # Project Context
    project_id = request.args.get("project_id", type=int)

    if not project_id:
        return jsonify({"success": False, "message": "Project ID required"}), 400

    conn = get_db_connection()
    try:
        with conn.cursor() as cursor:
            if request.method == 'GET':
                cursor.execute("SELECT * FROM project_configs WHERE project_id = %s", (project_id,))
                config = cursor.fetchone()
                if config:
                    return jsonify({
                        "success": True,
                        "frame_data": config["frame_data"],
                        "config_json": json.loads(config["config_json"]) if config["config_json"] else {}
                    })
                return jsonify({"success": True, "frame_data": None, "config_json": {}})
            
            elif request.method == 'POST':
                data = request.get_json(silent=True) or {}
                frame_data = data.get("frame_data") # base64 string or /uploads path
                try:
                    if frame_data:
                        frame_data = save_data_url(frame_data, "id-card-frame.png", "frames", project_id)
                except ValueError as err:
                    return jsonify({"success": False, "message": str(err)}), 400
                config_json = data.get("config_json") # dict

                config_str = json.dumps(config_json) if config_json is not None else None

                cursor.execute("""
                    INSERT INTO project_configs (project_id, frame_data, config_json)
                    VALUES (%s, %s, %s)
                    ON DUPLICATE KEY UPDATE
                        frame_data = IF(%s IS NOT NULL, %s, frame_data),
                        config_json = IF(%s IS NOT NULL, %s, config_json)
                """, (project_id, frame_data, config_str, frame_data, frame_data, config_str, config_str))
                conn.commit()
                return jsonify({
                    "success": True,
                    "message": "Configuration successfully saved",
                    "frame_data": frame_data
                })
    except Exception as e:
        print("Config API error:", e)
        return jsonify({"success": False, "message": "Internal Server Error"}), 500
    finally:
        conn.close()

# API: Student Records (Admin, Staff, or Student context)
@app.route('/api/students', methods=['GET'])
def api_get_students():
    session = get_current_session()
    if not session:
        return jsonify({"success": False, "message": "Session expired or invalid"}), 401

    search = request.args.get("search", "").strip()
    sort_key = request.args.get("sort", "id").strip().lower()
    direction = request.args.get("direction", "asc").strip().lower()
    direction = "DESC" if direction == "desc" else "ASC"
    include_archived = parse_bool_arg("include_archived", False)
    fetch_all = parse_bool_arg("all", False)
    page = max(1, request.args.get("page", default=1, type=int))
    page_size = max(1, min(request.args.get("page_size", default=50, type=int), 500))
    offset = (page - 1) * page_size

    conn = get_db_connection()
    try:
        with conn.cursor() as cursor:
            # 1. Student Role - fetch their own record only
            if session["role"] == "student":
                cursor.execute(
                    "SELECT * FROM students WHERE id = %s AND project_id = %s AND status <> 'archived'",
                    (session["student_id"], session["project_id"])
                )
                student = cursor.fetchone()
                return jsonify({
                    "success": True,
                    "students": [student] if student else [],
                    "total": 1 if student else 0,
                    "page": 1,
                    "page_size": 1,
                    "total_pages": 1 if student else 0
                })
            
            if session["role"] == "staff":
                project_id = session["project_id"]
                include_archived = False
            elif session["role"] == "admin":
                project_id = request.args.get("project_id", type=int)
                if not project_id:
                    return jsonify({"success": False, "message": "Project ID required"}), 400
            else:
                return jsonify({"success": False, "message": "Unauthorized"}), 403

            where_clauses = ["project_id = %s"]
            params = [project_id]
            if not include_archived:
                where_clauses.append("status <> 'archived'")

            if search:
                like_value = f"%{search}%"
                searchable_columns = [
                    "id", "roll_number", "name", "class", "section", "father_name",
                    "mother_name", "contact", "dob", "blood_group", "address", "status"
                ]
                where_clauses.append("(" + " OR ".join(f"{col} LIKE %s" for col in searchable_columns) + ")")
                params.extend([like_value] * len(searchable_columns))

            where_sql = " AND ".join(where_clauses)
            sort_column = STUDENT_SORT_COLUMNS.get(sort_key, "id")
            if sort_key == "photo":
                order_sql = f"ORDER BY CASE WHEN photo_data IS NULL OR photo_data = '' THEN 0 ELSE 1 END {direction}, id ASC"
            else:
                order_sql = f"ORDER BY {sort_column} {direction}, id ASC"

            cursor.execute(f"SELECT COUNT(*) as cnt FROM students WHERE {where_sql}", tuple(params))
            total = cursor.fetchone()["cnt"]

            if fetch_all:
                cursor.execute(f"SELECT * FROM students WHERE {where_sql} {order_sql}", tuple(params))
                students = cursor.fetchall()
                page_size = len(students) or total or 1
                page = 1
            else:
                cursor.execute(
                    f"SELECT * FROM students WHERE {where_sql} {order_sql} LIMIT %s OFFSET %s",
                    tuple(params + [page_size, offset])
                )
                students = cursor.fetchall()

            total_pages = (total + page_size - 1) // page_size if total else 0
            return jsonify({
                "success": True,
                "students": students,
                "total": total,
                "page": page,
                "page_size": page_size,
                "total_pages": total_pages,
                "sort": sort_key,
                "direction": direction.lower(),
                "search": search,
                "include_archived": include_archived
            })

    except Exception as e:
        print("Fetch students error:", e)
        return jsonify({"success": False, "message": "Internal Server Error"}), 500
    finally:
        conn.close()

# API: Update Student Record (With Edit Count limits for Students)
@app.route('/api/students/<student_id>', methods=['PUT'])
def api_update_student(student_id):
    session = get_current_session()
    if not session:
        return jsonify({"success": False, "message": "Unauthorized"}), 401

    # Security check:
    # Students can only update their own record and within project context
    if session["role"] == "student" and session["student_id"] != student_id:
        return jsonify({"success": False, "message": "Forbidden"}), 403

    project_id = session["project_id"]
    if session["role"] == "admin":
        project_id = request.args.get("project_id", type=int)
        if not project_id:
            return jsonify({"success": False, "message": "Project ID required for admin edit"}), 400

    data = request.get_json(silent=True) or {}
    if session["role"] == "staff" and "photo_data" in data:
        return jsonify({"success": False, "message": "Unauthorized"}), 403
    
    conn = get_db_connection()
    try:
        with conn.cursor() as cursor:
            # Check existing student profile
            cursor.execute("SELECT * FROM students WHERE id = %s AND project_id = %s", (student_id, project_id))
            student = cursor.fetchone()
            if not student:
                return jsonify({"success": False, "message": "Student profile not found!"}), 404

            # Enforce 2-edit limit if role is student
            if session["role"] == "student":
                if student["edit_count"] >= 2:
                    return jsonify({
                        "success": False,
                        "message": "Maximum edit limit reached! You are allowed to edit your profile exactly twice."
                    }), 400
                
                # Student cannot modify 'status' or 'edit_count' manually
                status = student["status"]
                new_edit_count = student["edit_count"] + 1
            else:
                # Admin/Staff can change status and doesn't increment edit_count
                status = data.get("status", student["status"])
                new_edit_count = student["edit_count"]

            # Read edit fields (fallback to existing values if not supplied)
            roll_number = data.get("roll_number", student["roll_number"])
            name = data.get("name", student["name"])
            class_val = data.get("class", student["class"])
            section = data.get("section", student["section"])
            father_name = data.get("father_name", student["father_name"])
            mother_name = data.get("mother_name", student["mother_name"])
            contact = data.get("contact", student["contact"])
            dob = data.get("dob", student["dob"])
            blood_group = data.get("blood_group", student["blood_group"])
            address = data.get("address", student["address"])
            correction_note = data.get("correction_note", student.get("correction_note"))
            photo_data = data.get("photo_data", student["photo_data"])
            try:
                if photo_data:
                    photo_data = save_data_url(photo_data, f"{student_id}.jpg", "students", project_id)
            except ValueError as err:
                return jsonify({"success": False, "message": str(err)}), 400

            cursor.execute("""
                UPDATE students SET
                    roll_number = %s,
                    name = %s,
                    class = %s,
                    section = %s,
                    father_name = %s,
                    mother_name = %s,
                    contact = %s,
                    dob = %s,
                    blood_group = %s,
                    address = %s,
                    correction_note = %s,
                    status = %s,
                    edit_count = %s,
                    photo_data = %s
                WHERE id = %s AND project_id = %s
            """, (roll_number, name, class_val, section, father_name, mother_name,
                  contact, dob, blood_group, address, correction_note, status, new_edit_count, photo_data,
                  student_id, project_id))
            
            conn.commit()
            return jsonify({
                "success": True,
                "message": f"Successfully updated student {name}",
                "edit_count": new_edit_count
            })
    except Exception as e:
        print("Update student error:", e)
        return jsonify({"success": False, "message": "Internal Server Error"}), 500
    finally:
        conn.close()


@app.route('/api/students/<student_id>', methods=['DELETE'])
def api_delete_student(student_id):
    session = get_current_session()
    if not session:
        return jsonify({"success": False, "message": "Session expired or invalid"}), 401
    if session["role"] != "admin":
        return jsonify({"success": False, "message": "Unauthorized"}), 403

    project_id = request.args.get("project_id", type=int)
    if not project_id:
        return jsonify({"success": False, "message": "Project ID required"}), 400

    conn = get_db_connection()
    try:
        with conn.cursor() as cursor:
            cursor.execute("DELETE FROM students WHERE id = %s AND project_id = %s", (student_id, project_id))
        conn.commit()
        return jsonify({"success": True, "message": "Student deleted successfully"})
    except Exception as e:
        print("Delete student error:", e)
        return jsonify({"success": False, "message": "Internal Server Error"}), 500
    finally:
        conn.close()


@app.route('/api/students/<student_id>/archive', methods=['POST'])
def api_archive_student(student_id):
    session = get_current_session()
    if not session:
        return jsonify({"success": False, "message": "Session expired or invalid"}), 401
    if session["role"] != "admin":
        return jsonify({"success": False, "message": "Unauthorized"}), 403

    project_id = request.args.get("project_id", type=int)
    if not project_id:
        return jsonify({"success": False, "message": "Project ID required"}), 400

    data = request.get_json(silent=True) or {}
    archived = data.get("archived", True)
    conn = get_db_connection()
    try:
        with conn.cursor() as cursor:
            if archived:
                cursor.execute("UPDATE students SET status = 'archived' WHERE id = %s AND project_id = %s", (student_id, project_id))
                message = "Student archived successfully"
            else:
                cursor.execute("UPDATE students SET status = 'pending' WHERE id = %s AND project_id = %s", (student_id, project_id))
                message = "Student restored successfully"
        conn.commit()
        return jsonify({"success": True, "message": message, "status": "archived" if archived else "pending"})
    except Exception as e:
        print("Archive student error:", e)
        return jsonify({"success": False, "message": "Internal Server Error"}), 500
    finally:
        conn.close()


# API: Student Self-Verification (student marks own profile as verified, no edit count used)
@app.route('/api/students/<student_id>/verify', methods=['POST'])
def api_student_verify(student_id):
    session = get_current_session()
    if not session:
        return jsonify({"success": False, "message": "Session expired or invalid"}), 401
    if session["role"] != "student":
        return jsonify({"success": False, "message": "Unauthorized"}), 403
    # Students can only verify their own profile
    if session.get("student_id") != student_id:
        return jsonify({"success": False, "message": "Forbidden"}), 403

    project_id = session["project_id"]
    data = request.get_json(silent=True) or {}
    requested_status = data.get("status", "verified")
    # Students can only toggle between pending and verified, not archive themselves
    if requested_status not in ("pending", "verified"):
        requested_status = "verified"

    conn = get_db_connection()
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                "UPDATE students SET status = %s WHERE id = %s AND project_id = %s AND status != 'archived'",
                (requested_status, student_id, project_id)
            )
        conn.commit()
        return jsonify({"success": True, "message": f"Status updated to {requested_status}", "status": requested_status})
    except Exception as e:
        print("Student verify error:", e)
        return jsonify({"success": False, "message": "Could not update verification status."}), 500
    finally:
        conn.close()


# API: Student Correction Notes (Staff or Admin)
@app.route('/api/students/notes', methods=['GET'])
def api_student_notes():
    session = get_current_session()
    if not session:
        return jsonify({"success": False, "message": "Session expired or invalid"}), 401
    if session["role"] not in ("admin", "staff"):
        return jsonify({"success": False, "message": "Unauthorized"}), 403

    if session["role"] == "staff":
        project_id = session["project_id"]
    else:
        project_id = request.args.get("project_id", type=int)
        if not project_id:
            return jsonify({"success": False, "message": "Project ID required"}), 400

    conn = get_db_connection()
    try:
        with conn.cursor() as cursor:
            cursor.execute("""
                SELECT id, name, class, section, correction_note
                FROM students
                WHERE project_id = %s
                  AND correction_note IS NOT NULL
                  AND correction_note != ''
                  AND status != 'archived'
                ORDER BY id ASC
            """, (project_id,))
            notes = cursor.fetchall()
        return jsonify({"success": True, "notes": notes})
    except Exception as e:
        print("Student notes API error:", e)
        return jsonify({"success": False, "message": "Could not load correction notes."}), 500
    finally:
        conn.close()


@app.route('/api/students/<student_id>/note', methods=['DELETE'])
def api_delete_student_note(student_id):
    session = get_current_session()
    if not session:
        return jsonify({"success": False, "message": "Session expired or invalid"}), 401
    if session["role"] not in ("admin", "staff"):
        return jsonify({"success": False, "message": "Unauthorized"}), 403

    if session["role"] == "staff":
        project_id = session["project_id"]
    else:
        project_id = request.args.get("project_id", type=int)
        if not project_id:
            return jsonify({"success": False, "message": "Project ID required"}), 400

    conn = get_db_connection()
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                "UPDATE students SET correction_note = NULL WHERE id = %s AND project_id = %s",
                (student_id, project_id)
            )
        conn.commit()
        return jsonify({"success": True, "message": "Correction note cleared successfully"})
    except Exception as e:
        print("Delete student note error:", e)
        return jsonify({"success": False, "message": "Could not clear note."}), 500
    finally:
        conn.close()


# API: Excel & Photo Import matching (Admin only)
@app.route('/api/students/upload', methods=['POST'])
def api_upload_data():
    session = get_current_session()
    if not session:
        return jsonify({"success": False, "message": "Session expired or invalid"}), 401
    if session["role"] != "admin":
        return jsonify({"success": False, "message": "Unauthorized"}), 403

    project_id = request.form.get("project_id", type=int)

    if not project_id:
        return jsonify({"success": False, "message": "Project ID required"}), 400

    excel_file = request.files.get("excel")
    photos = request.files.getlist("photos")
    if not excel_file and not photos:
        return jsonify({"success": False, "message": "Please upload an Excel sheet or student photos."}), 400

    conn = get_db_connection()
    try:
        parsed_students = []
        
        # 1. Parse Excel file if present
        if excel_file:
            excel_ext = get_file_extension(excel_file.filename)
            if excel_ext not in ALLOWED_EXCEL_EXTENSIONS:
                return jsonify({"success": False, "message": "Please upload an .xlsx Excel file."}), 400
            if not looks_like_xlsx(excel_file):
                return jsonify({"success": False, "message": "The Excel file is not a valid .xlsx workbook."}), 400

            from openpyxl import load_workbook
            in_mem_file = io.BytesIO(excel_file.read())
            wb = load_workbook(in_mem_file, read_only=True, data_only=True)
            sheet = wb.active
            
            rows = list(sheet.iter_rows(values_only=True))
            if len(rows) > 0:
                def resolve_sheet_headers(row):
                    headers = [str(cell).strip() if cell is not None else "" for cell in row]
                    photo_idx = find_header_index(headers, [
                        'photo id', 'photo no', 'photo number', 'image id', 'picture id',
                        'file name', 'filename', 'photo', 'pic', 'picture', 'image'
                    ])
                    id_idx = find_header_index(headers, [
                        'student id', 'student code', 'registration no', 'registration number',
                        'admission no', 'admission number', 'reg no', 'scholar no', 'id'
                    ], excluded_indexes=[photo_idx])
                    roll_idx = find_header_index(headers, ['roll number', 'roll no', 'roll'])
                    name_idx = find_header_index(headers, ['student name', 'name', 'full name', 'candidate name'])
                    class_section_idx = find_header_index(headers, ['class and section', 'class section', 'class & section', 'class-section'])
                    class_idx = find_header_index(headers, ['class', 'grade', 'standard'])
                    sec_idx = find_header_index(headers, ['section', 'sec', 'batch'])
                    if class_idx == class_section_idx:
                        class_idx = -1
                    if sec_idx == class_section_idx:
                        sec_idx = -1
                    return {
                        "photo": photo_idx,
                        "id": id_idx,
                        "roll": roll_idx,
                        "name": name_idx,
                        "class_section": class_section_idx,
                        "class": class_idx,
                        "section": sec_idx,
                        "father": find_header_index(headers, ['father name', 'father', "father's name"]),
                        "mother": find_header_index(headers, ['mother name', 'mother', "mother's name"]),
                        "contact": find_header_index(headers, ['phonenumber', 'phone number', 'contact', 'mobile', 'phone', 'cell', 'mobile no']),
                        "dob": find_header_index(headers, ['dob', 'date of birth', 'birth date']),
                        "blood": find_header_index(headers, ['blood group', 'blood', 'bg']),
                        "address": find_header_index(headers, ['address', 'residence', 'location']),
                    }

                def sheet_header_score(indexes):
                    if indexes["id"] == -1 and indexes["photo"] == -1:
                        return 0
                    return sum(1 for value in indexes.values() if value != -1)

                best_header_row_idx = 0
                best_indexes = resolve_sheet_headers(rows[0])
                best_score = sheet_header_score(best_indexes)
                if best_score == 0:
                    best_indexes = {key: -1 for key in best_indexes}
                for idx, candidate_row in enumerate(rows[:25]):
                    candidate_indexes = resolve_sheet_headers(candidate_row)
                    candidate_score = sheet_header_score(candidate_indexes)
                    if candidate_score > best_score:
                        best_header_row_idx = idx
                        best_indexes = candidate_indexes
                        best_score = candidate_score

                first_data_row_idx = best_header_row_idx + 1
                photo_idx = best_indexes["photo"]
                id_idx = best_indexes["id"]
                roll_idx = best_indexes["roll"]
                name_idx = best_indexes["name"]
                class_section_idx = best_indexes["class_section"]
                class_idx = best_indexes["class"]
                sec_idx = best_indexes["section"]
                father_idx = best_indexes["father"]
                mother_idx = best_indexes["mother"]
                contact_idx = best_indexes["contact"]
                dob_idx = best_indexes["dob"]
                blood_idx = best_indexes["blood"]
                address_idx = best_indexes["address"]

                detected_indexes = [
                    photo_idx, id_idx, roll_idx, name_idx, class_section_idx, class_idx, sec_idx,
                    father_idx, mother_idx, contact_idx, dob_idx, blood_idx, address_idx
                ]
                if all(idx == -1 for idx in detected_indexes):
                    first_non_empty_idx = -1
                    for idx, row in enumerate(rows):
                        if row and any(cell is not None and str(cell).strip() for cell in row):
                            first_non_empty_idx = idx
                            break
                    if first_non_empty_idx != -1:
                        first_value_idx = next(
                            (
                                idx for idx, cell in enumerate(rows[first_non_empty_idx])
                                if cell is not None and str(cell).strip()
                            ),
                            -1
                        )
                        if first_value_idx != -1 and looks_like_photo_reference(rows[first_non_empty_idx][first_value_idx]):
                            first_data_row_idx = first_non_empty_idx
                            id_idx = first_value_idx

                for row in rows[first_data_row_idx:]:
                    if not row or all(cell is None for cell in row):
                        continue

                    def cell_val(idx):
                        if idx != -1 and idx < len(row) and row[idx] is not None:
                            return str(row[idx]).strip()
                        return ""

                    photo_id = cell_val(photo_idx)
                    s_id = cell_val(id_idx) or photo_id
                    if not s_id or s_id.upper() == 'UNKNOWN':
                        continue

                    combined_class, combined_section = split_class_section(cell_val(class_section_idx))
                    class_value = cell_val(class_idx) or combined_class
                    section_value = cell_val(sec_idx) or combined_section

                    parsed_students.append({
                        "id": s_id,
                        "photo_id": photo_id or s_id,
                        "roll_number": cell_val(roll_idx),
                        "name": cell_val(name_idx) or "Unknown Name",
                        "class": class_value,
                        "section": section_value,
                        "father_name": cell_val(father_idx),
                        "mother_name": cell_val(mother_idx),
                        "contact": cell_val(contact_idx),
                        "dob": cell_val(dob_idx),
                        "blood_group": cell_val(blood_idx),
                        "address": cell_val(address_idx),
                        "status": "pending",
                        "photo_data": None
                    })

        # 2. Process Photos if present
        photo_candidates = {}
        photo_entries = []
        matched_photo_paths = set()
        if photos and len(photos) > 0:
            for photo in photos:
                if photo and photo.filename:
                    ext = get_file_extension(photo.filename)
                    if ext not in ALLOWED_IMAGE_EXTENSIONS:
                        return jsonify({"success": False, "message": f"Unsupported image type: {photo.filename}"}), 400
                    if not looks_like_image(photo):
                        return jsonify({"success": False, "message": f"Invalid image file: {photo.filename}"}), 400

            for photo in photos:
                if photo and photo.filename:
                    saved_path = save_upload_file(photo, "students", project_id)
                    photo_entry = {
                        "filename": photo.filename,
                        "path": saved_path
                    }
                    photo_entries.append(photo_entry)
                    add_photo_candidate(photo_candidates, photo.filename, photo_entry)

        # 3. Synchronize database records
        # If excel is provided, we merge records. If not, we just update matching photos on existing students.
        with conn.cursor() as cursor:
            # If Excel is uploaded, insert/update database records
            if parsed_students:
                for s in parsed_students:
                    # photo_id already falls back to the record ID when no photo column value exists.
                    photo_entry = find_unique_photo_match(photo_candidates, s["photo_id"])
                    if photo_entry and photo_entry["path"] in matched_photo_paths:
                        photo_entry = None
                    matched_photo = photo_entry["path"] if photo_entry else None
                    if photo_entry:
                        matched_photo_paths.add(matched_photo)

                    cursor.execute("""
                        INSERT INTO students (
                            id, project_id, photo_id, roll_number, name, class, section,
                            father_name, mother_name, contact, dob, blood_group, address,
                            status, edit_count, photo_data
                        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'pending', 0, %s)
                        ON DUPLICATE KEY UPDATE
                            photo_id = VALUES(photo_id),
                            roll_number = VALUES(roll_number),
                            name = VALUES(name),
                            class = VALUES(class),
                            section = VALUES(section),
                            father_name = VALUES(father_name),
                            mother_name = VALUES(mother_name),
                            contact = VALUES(contact),
                            dob = VALUES(dob),
                            blood_group = VALUES(blood_group),
                            address = VALUES(address),
                            photo_data = IF(%s IS NOT NULL, %s, photo_data)
                    """, (s["id"], project_id, s["photo_id"], s["roll_number"], s["name"], s["class"], s["section"],
                          s["father_name"], s["mother_name"], s["contact"], s["dob"], s["blood_group"], s["address"],
                          matched_photo, matched_photo, matched_photo))
            
            # If ONLY photos were uploaded, match them to current students in this project
            elif photo_candidates:
                cursor.execute("SELECT id, photo_id FROM students WHERE project_id = %s", (project_id,))
                existing_students = cursor.fetchall()
                for es in existing_students:
                    sid = es["id"]
                    photo_entry = find_unique_photo_match(photo_candidates, es.get("photo_id") or sid)
                    if photo_entry and photo_entry["path"] in matched_photo_paths:
                        photo_entry = None
                    matched_photo = photo_entry["path"] if photo_entry else None
                    if photo_entry:
                        matched_photo_paths.add(matched_photo)
                        cursor.execute(
                            "UPDATE students SET photo_data = %s WHERE id = %s AND project_id = %s",
                            (matched_photo, sid, project_id)
                        )
            
            conn.commit()

        # Count total students now
        with conn.cursor() as cursor:
            cursor.execute("SELECT COUNT(*) as cnt FROM students WHERE project_id = %s", (project_id,))
            tot = cursor.fetchone()["cnt"]
            cursor.execute("SELECT COUNT(*) as cnt FROM students WHERE project_id = %s AND photo_data IS NOT NULL", (project_id,))
            matched = cursor.fetchone()["cnt"]

        return jsonify({
            "success": True,
            "message": "Data processed and stored successfully",
            "summary": {
                "total_records": tot,
                "photos_matched": matched,
                "uploaded_photo_count": len(photo_entries),
                "matched_photo_files": [
                    entry["filename"] for entry in photo_entries if entry["path"] in matched_photo_paths
                ],
                "unmatched_photos": [
                    entry["filename"] for entry in photo_entries if entry["path"] not in matched_photo_paths
                ]
            }
        })
    except Exception as e:
        print("Upload API error:", e)
        return jsonify({"success": False, "message": "Failed to parse files"}), 500
    finally:
        conn.close()

if __name__ == '__main__':
    # Pterodactyl passes the allocation through SERVER_PORT; PORT is the local fallback.
    port = int(os.getenv("SERVER_PORT", os.getenv("PORT", 5000)))
    print(f"Starting Satish Ad Agency web portal on 0.0.0.0:{port}...")
    app.run(host='0.0.0.0', port=port, debug=False)
