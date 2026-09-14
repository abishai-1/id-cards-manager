import base64
import io
import importlib
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook


os.environ["SKIP_DB_INIT"] = "true"

PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAAC0lEQVR4nGNgYAAAAAMAASsJTYQAAAAASUVORK5CYII="
)


class FakeCursor:
    def __init__(self, connection):
        self.connection = connection

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, tb):
        return False

    def execute(self, sql, params=None):
        self.connection.executed.append((sql, params))
        if "SELECT COUNT(*) as cnt" in sql:
            self.connection.pending_fetches.append({"cnt": 1})
        elif "SELECT * FROM projects WHERE access_code" in sql:
            self.connection.pending_fetches.append({
                "id": 77,
                "name": "Test School",
                "access_code": "SAT-ABC12345",
            })
        elif "SELECT id, name, access_code FROM projects WHERE id" in sql:
            self.connection.pending_fetches.append({
                "id": 77,
                "name": "Test School",
                "access_code": "SAT-NEWCODE",
            })
        elif "INNER JOIN projects p ON p.id = s.project_id" in sql:
            self.connection.pending_fetches.append([{
                "id": "CODX-001",
                "photo_id": "IMG-001",
                "project_id": 77,
                "roll_number": "1",
                "name": "Codex Test Student",
                "class": "10",
                "section": "A",
                "status": "pending",
            }])
        elif "SELECT * FROM students WHERE id = %s" in sql:
            self.connection.pending_fetches.append({
                "id": "CODX-001",
                "photo_id": "IMG-001",
                "project_id": 77,
                "roll_number": "1",
                "name": "Codex Test Student",
                "class": "10",
                "section": "A",
                "father_name": "Test Father",
                "mother_name": "Test Mother",
                "contact": "9876543210",
                "dob": "2012-05-24",
                "blood_group": "O+",
                "address": "123 Test Street",
                "correction_note": "",
                "status": "pending",
                "edit_count": 0,
                "photo_data": None,
            })
        elif "SELECT * FROM project_configs" in sql:
            self.connection.pending_fetches.append(None)
        elif "SHOW COLUMNS FROM students LIKE 'correction_note'" in sql:
            self.connection.pending_fetches.append({"Field": "correction_note"})

    def fetchone(self):
        if self.connection.pending_fetches:
            return self.connection.pending_fetches.pop(0)
        return None

    def fetchall(self):
        if self.connection.pending_fetches:
            value = self.connection.pending_fetches.pop(0)
            return value if isinstance(value, list) else [value] if value else []
        if any("contact_messages" in sql for sql, _ in self.connection.executed[-1:]):
            return [{"id": 1, "name": "Codex Test", "email": "codex@example.com", "message": "Hello", "created_at": "2026-05-24"}]
        return [{"id": "CODX-001"}]


class FakeConnection:
    def __init__(self):
        self.executed = []
        self.pending_fetches = []
        self.commits = 0

    def cursor(self):
        return FakeCursor(self)

    def commit(self):
        self.commits += 1

    def close(self):
        pass


class LoginCandidateConnection(FakeConnection):
    def __init__(self, candidates):
        super().__init__()
        self.candidates = candidates

    def cursor(self):
        return LoginCandidateCursor(self)


class LoginCandidateCursor(FakeCursor):
    def execute(self, sql, params=None):
        self.connection.executed.append((sql, params))
        if "INNER JOIN projects p ON p.id = s.project_id" in sql:
            self.connection.pending_fetches.append(self.connection.candidates)


class BackendWorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        os.environ.setdefault("DB_HOST", "127.0.0.1")
        cls.main = importlib.import_module("main")

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.upload_dir = Path(self.tmp.name) / "uploads"
        self.main.app.config["TESTING"] = True
        self.main.app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024
        self.client = self.main.app.test_client()

    def tearDown(self):
        self.tmp.cleanup()

    def make_workbook(self, headers=None, row=None):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers or ["pic", "name", "address", "phonenumber", "mother name", "father name", "dob", "class and section", "blood group"])
        sheet.append(row or ["CODX-001", "Codex Test Student", "123 Test Street", "9876543210", "Test Mother", "Test Father", "2012-05-24", "10-A", "O+"])
        data = io.BytesIO()
        workbook.save(data)
        data.seek(0)
        return data

    def test_excel_photo_upload_saves_photo_to_uploads_and_parses_sheet(self):
        fake_db = FakeConnection()
        with patch.object(self.main, "UPLOAD_FOLDER", str(self.upload_dir)), \
             patch.object(self.main, "get_current_session", return_value={"role": "admin"}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.post(
                "/api/students/upload",
                data={
                    "project_id": "77",
                    "excel": (self.make_workbook(), "students.xlsx"),
                    "photos": (io.BytesIO(PNG_BYTES), "CODX-001.png"),
                },
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        payload = response.get_json()
        self.assertTrue(payload["success"])
        inserted = [params for sql, params in fake_db.executed if "INSERT INTO students" in sql][0]
        self.assertEqual(inserted[0], "CODX-001")
        self.assertEqual(inserted[2], "CODX-001")
        self.assertEqual(inserted[4], "Codex Test Student")
        self.assertEqual(inserted[5], "10")
        self.assertEqual(inserted[6], "A")
        self.assertEqual(inserted[9], "9876543210")
        self.assertTrue(inserted[13].startswith("/uploads/students/77/"), inserted[13])
        saved_path = Path(str(self.upload_dir)) / "students" / "77"
        self.assertEqual(len(list(saved_path.glob("*.png"))), 1)

    def test_frame_config_saves_frame_to_uploads(self):
        fake_db = FakeConnection()
        frame_data = "data:image/png;base64," + base64.b64encode(PNG_BYTES).decode("ascii")
        with patch.object(self.main, "UPLOAD_FOLDER", str(self.upload_dir)), \
             patch.object(self.main, "get_current_session", return_value={"role": "admin"}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.post(
                "/api/config?project_id=77",
                json={"frame_data": frame_data, "config_json": {"photo": {"show": True}}},
            )

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        payload = response.get_json()
        self.assertTrue(payload["success"])
        self.assertTrue(payload["frame_data"].startswith("/uploads/frames/77/"))
        self.assertEqual(len(list((self.upload_dir / "frames" / "77").glob("*.png"))), 1)

    def test_staff_cannot_save_frame_config_even_with_direct_api_request(self):
        fake_db = FakeConnection()
        frame_data = "data:image/png;base64," + base64.b64encode(PNG_BYTES).decode("ascii")
        with patch.object(self.main, "UPLOAD_FOLDER", str(self.upload_dir)), \
             patch.object(self.main, "get_current_session", return_value={"role": "staff", "project_id": 77}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.post(
                "/api/config?project_id=99",
                json={"frame_data": frame_data, "config_json": {"photo": {"show": True}}},
            )

        self.assertEqual(response.status_code, 403)
        self.assertFalse(response.get_json()["success"])
        self.assertFalse(any("INSERT INTO project_configs" in sql for sql, _ in fake_db.executed))
        self.assertFalse((self.upload_dir / "frames" / "77").exists())

    def test_staff_cannot_read_frame_config(self):
        fake_db = FakeConnection()
        with patch.object(self.main, "get_current_session", return_value={"role": "staff", "project_id": 77}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.get("/api/config?project_id=99")

        self.assertEqual(response.status_code, 403)
        self.assertFalse(response.get_json()["success"])
        self.assertFalse(any("project_configs" in sql for sql, _ in fake_db.executed))

    def test_invalid_frame_data_is_rejected_without_saving_file(self):
        fake_db = FakeConnection()
        with patch.object(self.main, "UPLOAD_FOLDER", str(self.upload_dir)), \
             patch.object(self.main, "get_current_session", return_value={"role": "admin"}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.post(
                "/api/config?project_id=77",
                json={"frame_data": "data:image/svg+xml;base64,PHN2Zz48L3N2Zz4=", "config_json": {}},
            )

        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.get_json()["success"])
        self.assertFalse((self.upload_dir / "frames" / "77").exists())

    def test_public_catalog_returns_saved_image_overrides(self):
        fake_db = FakeConnection()
        expected_images = [f"/uploads/catalog/id-cards/image-{index}.png" for index in range(1, 5)]
        fake_db.pending_fetches.append([{
            "item_key": "id-cards",
            "images_json": json.dumps(expected_images),
        }])
        with patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.get("/api/catalog")

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertEqual(response.get_json()["catalog"]["id-cards"], expected_images)

    def test_admin_can_replace_catalog_gallery_with_exactly_four_images(self):
        fake_db = FakeConnection()
        with patch.object(self.main, "UPLOAD_FOLDER", str(self.upload_dir)), \
             patch.object(self.main, "get_current_session", return_value={"role": "admin"}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.open(
                "/api/catalog/id-cards",
                method="PUT",
                data={
                    "images": [
                        (io.BytesIO(PNG_BYTES), "one.png"),
                        (io.BytesIO(PNG_BYTES), "two.png"),
                        (io.BytesIO(PNG_BYTES), "three.png"),
                        (io.BytesIO(PNG_BYTES), "four.png"),
                    ],
                },
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        payload = response.get_json()
        self.assertEqual(len(payload["images"]), 4)
        self.assertTrue(all(image.startswith("/uploads/catalog/id-cards/") for image in payload["images"]))
        self.assertEqual(len(list((self.upload_dir / "catalog" / "id-cards").glob("*.png"))), 4)
        self.assertTrue(any("INSERT INTO catalog_items" in sql for sql, _ in fake_db.executed))

    def test_staff_cannot_replace_catalog_gallery_through_direct_api_request(self):
        fake_db = FakeConnection()
        image_data = "data:image/png;base64," + base64.b64encode(PNG_BYTES).decode("ascii")
        with patch.object(self.main, "UPLOAD_FOLDER", str(self.upload_dir)), \
             patch.object(self.main, "get_current_session", return_value={"role": "staff", "project_id": 77}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.put(
                "/api/catalog/id-cards",
                json={"images": [image_data] * 4},
            )

        self.assertEqual(response.status_code, 403)
        self.assertFalse(response.get_json()["success"])
        self.assertFalse(any("INSERT INTO catalog_items" in sql for sql, _ in fake_db.executed))
        self.assertFalse((self.upload_dir / "catalog").exists())

    def test_catalog_gallery_rejects_anything_other_than_four_images(self):
        fake_db = FakeConnection()
        with patch.object(self.main, "UPLOAD_FOLDER", str(self.upload_dir)), \
             patch.object(self.main, "get_current_session", return_value={"role": "admin"}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.open(
                "/api/catalog/id-cards",
                method="PUT",
                data={
                    "images": [
                        (io.BytesIO(PNG_BYTES), "one.png"),
                        (io.BytesIO(PNG_BYTES), "two.png"),
                        (io.BytesIO(PNG_BYTES), "three.png"),
                    ],
                },
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.get_json()["success"])
        self.assertFalse((self.upload_dir / "catalog").exists())

    def test_catalog_gallery_still_accepts_existing_json_image_payloads(self):
        fake_db = FakeConnection()
        image_data = "data:image/png;base64," + base64.b64encode(PNG_BYTES).decode("ascii")
        with patch.object(self.main, "UPLOAD_FOLDER", str(self.upload_dir)), \
             patch.object(self.main, "get_current_session", return_value={"role": "admin"}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.put(
                "/api/catalog/id-cards",
                json={"images": [image_data] * 4},
            )

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertEqual(len(response.get_json()["images"]), 4)

    def test_contact_endpoint_stores_message(self):
        fake_db = FakeConnection()
        with patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.post(
                "/api/contact",
                json={"name": "Codex Test", "email": "codex@example.com", "message": "Hello"},
            )

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertTrue(response.get_json()["success"])
        self.assertTrue(any("INSERT INTO contact_messages" in sql for sql, _ in fake_db.executed))

    def test_private_project_files_are_not_served_by_catch_all_route(self):
        for path in ("/.env", "/main.py", "/tests/test_backend_workflow.py"):
            response = self.client.get(path)
            self.assertEqual(response.status_code, 404, path)

    def test_admin_login_fails_closed_when_password_is_not_configured(self):
        with patch.object(self.main, "ADMIN_PASSWORD", ""), \
             patch.object(self.main, "get_db_connection", side_effect=AssertionError("DB should not be touched")):
            response = self.client.post(
                "/api/auth/login",
                json={"role": "admin", "username": "satish", "password": ""},
            )

        self.assertEqual(response.status_code, 503)
        self.assertFalse(response.get_json()["success"])

    def test_login_returns_json_when_database_is_unavailable(self):
        with patch.object(self.main, "get_db_connection", side_effect=OSError("Database offline")):
            response = self.client.post(
                "/api/auth/login",
                json={"role": "admin", "username": "satish", "password": "incorrect"},
            )

        self.assertEqual(response.status_code, 503)
        self.assertTrue(response.is_json)
        self.assertFalse(response.get_json()["success"])
        self.assertIn("Database unavailable", response.get_json()["message"])

    def test_staff_login_uses_project_access_code(self):
        fake_db = FakeConnection()
        with patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.post(
                "/api/auth/login",
                json={"role": "staff", "code": "SAT-ABC12345"},
            )

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        payload = response.get_json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["userContext"]["project_id"], 77)
        self.assertTrue(any("WHERE access_code = %s" in sql for sql, _ in fake_db.executed))

    def test_student_login_is_scoped_by_school_class_section_roll_and_photo_id(self):
        fake_db = FakeConnection()
        with patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.post(
                "/api/auth/login",
                json={
                    "role": "student",
                    "school": "Test School",
                    "class": "10",
                    "section": "A",
                    "roll_number": "1",
                    "photo_id": " img 001.jpg ",
                },
            )

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        payload = response.get_json()
        self.assertEqual(payload["userContext"]["id"], "CODX-001")
        self.assertEqual(payload["userContext"]["photo_id"], "IMG-001")
        joined_sql = "\n".join(sql for sql, _ in fake_db.executed)
        self.assertIn("INNER JOIN projects", joined_sql)
        self.assertNotIn("LOWER(TRIM(s.roll_number))", joined_sql)

    def test_student_login_ignores_missing_master_fields_and_uses_remaining_details(self):
        fake_db = LoginCandidateConnection([{
            "id": "P-001",
            "photo_id": "",
            "project_id": 77,
            "roll_number": "",
            "name": "A.ALLU",
            "class": "",
            "section": "A",
            "status": "pending",
        }])
        with patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.post(
                "/api/auth/login",
                json={
                    "role": "student",
                    "school": "Test School",
                    "class": "1",
                    "section": "A",
                    "roll_number": "15",
                    "photo_id": "P-001.jpg",
                },
            )

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        payload = response.get_json()
        self.assertEqual(payload["userContext"]["id"], "P-001")
        self.assertEqual(payload["userContext"]["photo_id"], "P-001")

    def test_student_login_requires_at_least_one_detail_beyond_school(self):
        fake_db = FakeConnection()
        with patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.post(
                "/api/auth/login",
                json={"role": "student", "school": "Test School"},
            )

        self.assertEqual(response.status_code, 400)
        self.assertIn("at least one student detail", response.get_json()["message"])

    def test_student_login_reports_overlap_when_remaining_details_are_ambiguous(self):
        fake_db = LoginCandidateConnection([
            {
                "id": "P-001",
                "photo_id": "",
                "project_id": 77,
                "roll_number": "",
                "name": "A.ALLU",
                "class": "",
                "section": "A",
                "status": "pending",
            },
            {
                "id": "P-002",
                "photo_id": "",
                "project_id": 77,
                "roll_number": "",
                "name": "B.BABU",
                "class": "",
                "section": "A",
                "status": "pending",
            },
        ])
        with patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.post(
                "/api/auth/login",
                json={"role": "student", "school": "Test School", "section": "A"},
            )

        self.assertEqual(response.status_code, 409)
        self.assertIn("contact your school staff", response.get_json()["message"])

    def test_photo_matching_detects_picture_id_and_normalizes_symbols_spaces_and_extensions(self):
        fake_db = FakeConnection()
        workbook = self.make_workbook(
            headers=["Student ID", "Picture ID", "Name", "Roll Number", "Class", "Section"],
            row=["CODX-001", " IMG-001.png ", "Codex Test Student", "1", "10", "A"],
        )
        with patch.object(self.main, "UPLOAD_FOLDER", str(self.upload_dir)), \
             patch.object(self.main, "get_current_session", return_value={"role": "admin"}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.post(
                "/api/students/upload",
                data={
                    "project_id": "77",
                    "excel": (workbook, "students.xlsx"),
                    "photos": (io.BytesIO(PNG_BYTES), "img 001.jpg"),
                },
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        inserted = [params for sql, params in fake_db.executed if "INSERT INTO students" in sql][0]
        self.assertEqual(inserted[0], "CODX-001")
        self.assertEqual(inserted[2], "IMG-001.png")
        self.assertTrue(inserted[13].startswith("/uploads/students/77/"), inserted[13])

    def test_generic_id_column_stays_student_id_when_picture_id_is_present(self):
        fake_db = FakeConnection()
        workbook = self.make_workbook(
            headers=["ID", "Picture ID", "Name", "Roll Number", "Class", "Section"],
            row=["STUDENT-001", " IMG-001.png ", "Codex Test Student", "1", "10", "A"],
        )
        with patch.object(self.main, "UPLOAD_FOLDER", str(self.upload_dir)), \
             patch.object(self.main, "get_current_session", return_value={"role": "admin"}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.post(
                "/api/students/upload",
                data={
                    "project_id": "77",
                    "excel": (workbook, "students.xlsx"),
                    "photos": (io.BytesIO(PNG_BYTES), "img 001.jpg"),
                },
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        inserted = [params for sql, params in fake_db.executed if "INSERT INTO students" in sql][0]
        self.assertEqual(inserted[0], "STUDENT-001")
        self.assertEqual(inserted[2], "IMG-001.png")
        self.assertTrue(inserted[13].startswith("/uploads/students/77/"), inserted[13])

    def test_id_column_with_photo_filenames_imports_and_matches_photos(self):
        fake_db = FakeConnection()
        workbook = self.make_workbook(
            headers=["ID"],
            row=["P-001.jpg"],
        )
        with patch.object(self.main, "UPLOAD_FOLDER", str(self.upload_dir)), \
             patch.object(self.main, "get_current_session", return_value={"role": "admin"}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.post(
                "/api/students/upload",
                data={
                    "project_id": "77",
                    "excel": (workbook, "students.xlsx"),
                    "photos": (io.BytesIO(PNG_BYTES), "P-001.jpg"),
                },
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        inserted = [params for sql, params in fake_db.executed if "INSERT INTO students" in sql][0]
        self.assertEqual(inserted[0], "P-001.jpg")
        self.assertEqual(inserted[2], "P-001.jpg")
        self.assertEqual(inserted[4], "Unknown Name")
        self.assertTrue(inserted[13].startswith("/uploads/students/77/"), inserted[13])

    def test_school_sheet_id_values_without_extensions_match_photo_filenames(self):
        fake_db = FakeConnection()
        workbook = self.make_workbook(
            headers=["ID", "NAME", "CLASS", "SECTION", "CONTACT", "ADDRESS"],
            row=["P-001", "A.ALLU", "1", "A", "9876543210", "FLAT-101"],
        )
        with patch.object(self.main, "UPLOAD_FOLDER", str(self.upload_dir)), \
             patch.object(self.main, "get_current_session", return_value={"role": "admin"}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.post(
                "/api/students/upload",
                data={
                    "project_id": "77",
                    "excel": (workbook, "students.xlsx"),
                    "photos": (io.BytesIO(PNG_BYTES), "P-001.jpg"),
                },
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        inserted = [params for sql, params in fake_db.executed if "INSERT INTO students" in sql][0]
        self.assertEqual(inserted[0], "P-001")
        self.assertEqual(inserted[2], "P-001")
        self.assertEqual(inserted[4], "A.ALLU")
        self.assertEqual(inserted[5], "1")
        self.assertEqual(inserted[6], "A")
        self.assertEqual(inserted[9], "9876543210")
        self.assertEqual(inserted[12], "FLAT-101")
        self.assertTrue(inserted[13].startswith("/uploads/students/77/"), inserted[13])

    def test_offset_school_sheet_finds_header_row_and_columns(self):
        fake_db = FakeConnection()
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([])
        sheet.append(["", "", "School student list"])
        sheet.append(["", "", "ID", "NAME", "CLASS", "SECTION", "CONTACT", "ADDRESS"])
        sheet.append(["", "", "P-001", "A.ALLU", "1", "A", "9876543210", "FLAT-101"])
        data = io.BytesIO()
        workbook.save(data)
        data.seek(0)

        with patch.object(self.main, "UPLOAD_FOLDER", str(self.upload_dir)), \
             patch.object(self.main, "get_current_session", return_value={"role": "admin"}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.post(
                "/api/students/upload",
                data={
                    "project_id": "77",
                    "excel": (data, "students.xlsx"),
                    "photos": (io.BytesIO(PNG_BYTES), "P-001.jpg"),
                },
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        inserted = [params for sql, params in fake_db.executed if "INSERT INTO students" in sql][0]
        self.assertEqual(inserted[0], "P-001")
        self.assertEqual(inserted[4], "A.ALLU")
        self.assertEqual(inserted[5], "1")
        self.assertEqual(inserted[6], "A")
        self.assertEqual(inserted[9], "9876543210")
        self.assertEqual(inserted[12], "FLAT-101")
        self.assertTrue(inserted[13].startswith("/uploads/students/77/"), inserted[13])

    def test_headerless_photo_filename_column_imports_from_first_row(self):
        fake_db = FakeConnection()
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["P-001.jpg"])
        sheet.append(["P-002.jpg"])
        data = io.BytesIO()
        workbook.save(data)
        data.seek(0)

        with patch.object(self.main, "UPLOAD_FOLDER", str(self.upload_dir)), \
             patch.object(self.main, "get_current_session", return_value={"role": "admin"}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.post(
                "/api/students/upload",
                data={
                    "project_id": "77",
                    "excel": (data, "students.xlsx"),
                    "photos": (io.BytesIO(PNG_BYTES), "P-001.jpg"),
                },
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        inserts = [params for sql, params in fake_db.executed if "INSERT INTO students" in sql]
        self.assertEqual([params[0] for params in inserts], ["P-001.jpg", "P-002.jpg"])
        self.assertTrue(inserts[0][13].startswith("/uploads/students/77/"), inserts[0][13])
        self.assertIsNone(inserts[1][13])

    def test_picture_id_does_not_fall_back_to_student_id_for_photo_matching(self):
        fake_db = FakeConnection()
        workbook = self.make_workbook(
            headers=["Student ID", "Picture ID", "Name", "Roll Number", "Class", "Section"],
            row=["STUDENT-001", "IMG-001", "Codex Test Student", "1", "10", "A"],
        )
        with patch.object(self.main, "UPLOAD_FOLDER", str(self.upload_dir)), \
             patch.object(self.main, "get_current_session", return_value={"role": "admin"}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.post(
                "/api/students/upload",
                data={
                    "project_id": "77",
                    "excel": (workbook, "students.xlsx"),
                    "photos": (io.BytesIO(PNG_BYTES), "STUDENT-001.png"),
                },
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        inserted = [params for sql, params in fake_db.executed if "INSERT INTO students" in sql][0]
        self.assertIsNone(inserted[13])
        self.assertIn("STUDENT-001.png", response.get_json()["summary"]["unmatched_photos"])

    def test_photo_matching_does_not_guess_between_duplicate_normalized_filenames(self):
        candidates = {}
        self.main.add_photo_candidate(candidates, "IMG-001.png", {"path": "/first", "filename": "IMG-001.png"})
        self.main.add_photo_candidate(candidates, "img 001.jpg", {"path": "/second", "filename": "img 001.jpg"})

        self.assertIsNone(self.main.find_unique_photo_match(candidates, "IMG-001"))

    def test_numeric_photo_ids_match_filenames_with_leading_zeroes(self):
        self.assertTrue(
            set(self.main.identifier_keys("1")).intersection(self.main.identifier_keys("001.png"))
        )

    def test_students_endpoint_supports_search_sort_pagination_and_archive_filter(self):
        fake_db = FakeConnection()
        with patch.object(self.main, "get_current_session", return_value={"role": "admin"}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.get(
                "/api/students?project_id=77&search=codex&sort=name&direction=desc&page=2&page_size=25&include_archived=true"
            )

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        payload = response.get_json()
        self.assertEqual(payload["page"], 2)
        self.assertEqual(payload["page_size"], 25)
        self.assertIn("total", payload)
        self.assertIn("total_pages", payload)
        joined_sql = "\n".join(sql for sql, _ in fake_db.executed)
        self.assertIn("LIKE", joined_sql)
        self.assertIn("ORDER BY name DESC", joined_sql)
        self.assertIn("LIMIT %s OFFSET %s", joined_sql)

    def test_staff_students_endpoint_ignores_project_id_url_changes(self):
        fake_db = FakeConnection()
        with patch.object(self.main, "get_current_session", return_value={"role": "staff", "project_id": 77}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.get("/api/students?project_id=99&search=codex&include_archived=true")

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        first_count_query = next(params for sql, params in fake_db.executed if "SELECT COUNT(*) as cnt FROM students" in sql)
        first_count_sql = next(sql for sql, _ in fake_db.executed if "SELECT COUNT(*) as cnt FROM students" in sql)
        self.assertEqual(first_count_query[0], 77)
        self.assertNotIn(99, first_count_query)
        self.assertIn("status <> 'archived'", first_count_sql)

    def test_staff_cannot_upload_student_photo_through_edit_api(self):
        with patch.object(self.main, "get_current_session", return_value={"role": "staff", "project_id": 77}), \
             patch.object(self.main, "get_db_connection", side_effect=AssertionError("DB should not be touched")):
            response = self.client.put(
                "/api/students/CODX-001",
                json={"photo_data": "data:image/png;base64,ignored"},
            )

        self.assertEqual(response.status_code, 403)
        self.assertFalse(response.get_json()["success"])

    def test_upload_response_reports_unmatched_photos(self):
        fake_db = FakeConnection()
        with patch.object(self.main, "UPLOAD_FOLDER", str(self.upload_dir)), \
             patch.object(self.main, "get_current_session", return_value={"role": "admin"}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.post(
                "/api/students/upload",
                data={
                    "project_id": "77",
                    "excel": (self.make_workbook(), "students.xlsx"),
                    "photos": [
                        (io.BytesIO(PNG_BYTES), "CODX-001.png"),
                        (io.BytesIO(PNG_BYTES), "NOPE.png"),
                    ],
                },
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        summary = response.get_json()["summary"]
        self.assertIn("unmatched_photos", summary)
        self.assertIn("NOPE.png", summary["unmatched_photos"])
        self.assertIn("matched_photo_files", summary)

    def test_staff_cannot_upload_data_even_with_direct_api_request(self):
        fake_db = FakeConnection()
        with patch.object(self.main, "UPLOAD_FOLDER", str(self.upload_dir)), \
             patch.object(self.main, "get_current_session", return_value={"role": "staff", "project_id": 77}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            response = self.client.post(
                "/api/students/upload",
                data={
                    "project_id": "99",
                    "excel": (self.make_workbook(), "students.xlsx"),
                    "photos": (io.BytesIO(PNG_BYTES), "CODX-001.png"),
                },
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 403)
        self.assertFalse(response.get_json()["success"])
        self.assertFalse(any("INSERT INTO students" in sql for sql, _ in fake_db.executed))

    def test_admin_can_archive_and_delete_single_student(self):
        fake_db = FakeConnection()
        with patch.object(self.main, "get_current_session", return_value={"role": "admin"}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            archive_response = self.client.post("/api/students/CODX-001/archive?project_id=77")
            delete_response = self.client.delete("/api/students/CODX-001?project_id=77")

        self.assertEqual(archive_response.status_code, 200, archive_response.get_data(as_text=True))
        self.assertEqual(delete_response.status_code, 200, delete_response.get_data(as_text=True))
        joined_sql = "\n".join(sql for sql, _ in fake_db.executed)
        self.assertIn("UPDATE students SET status = 'archived'", joined_sql)
        self.assertIn("DELETE FROM students", joined_sql)

    def test_admin_can_read_contact_messages_and_regenerate_access_codes(self):
        fake_db = FakeConnection()
        with patch.object(self.main, "generate_access_code", return_value="SAT-NEWCODE"), \
             patch.object(self.main, "get_current_session", return_value={"role": "admin"}), \
             patch.object(self.main, "get_db_connection", return_value=fake_db):
            contact_response = self.client.get("/api/contact")
            access_response = self.client.post("/api/projects/77/access-code/regenerate")

        self.assertEqual(contact_response.status_code, 200, contact_response.get_data(as_text=True))
        self.assertEqual(access_response.status_code, 200, access_response.get_data(as_text=True))
        self.assertIn("messages", contact_response.get_json())
        self.assertEqual(access_response.get_json()["project"]["access_code"], "SAT-NEWCODE")
        joined_sql = "\n".join(sql for sql, _ in fake_db.executed)
        self.assertIn("SELECT id, name, email, message, created_at FROM contact_messages", joined_sql)
        self.assertIn("UPDATE projects SET access_code = %s", joined_sql)


if __name__ == "__main__":
    unittest.main()
